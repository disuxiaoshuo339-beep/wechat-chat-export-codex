from __future__ import annotations

import csv

import openpyxl

from build_index import _build_xlsx, _write_csv_fallback


VALUES = ['=HYPERLINK("https://example.test/","test")', "+1+1", "-1+1", "@SUM(1)", "\t=1", "\r=1"]


def _rows():
    return [{"序号": i, "客户显示名": value, "消息数": 1} for i, value in enumerate(VALUES, 1)]


def test_untrusted_index_names_remain_text_in_saved_xlsx(tmp_path):
    path = tmp_path / "index.xlsx"
    _build_xlsx({}, _rows(), path)
    workbook = openpyxl.load_workbook(path, data_only=False)
    try:
        sheet = workbook["客户索引"]
        for row_number, value in enumerate(VALUES, 2):
            cell = sheet.cell(row_number, 2)
            assert cell.data_type == "s"
            assert cell.value == "'" + value
            assert sheet.cell(row_number, 7).value == 1
    finally:
        workbook.close()


def test_csv_fallback_escapes_formula_prefixes(tmp_path):
    path = tmp_path / "index.csv"
    _write_csv_fallback(_rows(), path)
    with path.open(encoding="utf-8-sig", newline="") as stream:
        rows = list(csv.DictReader(stream))
    assert [row["客户显示名"] for row in rows] == ["'" + value for value in VALUES]
