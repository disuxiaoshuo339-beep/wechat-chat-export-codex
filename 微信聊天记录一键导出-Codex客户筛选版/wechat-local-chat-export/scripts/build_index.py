from __future__ import annotations

import argparse
import csv
import json
import sys
from pathlib import Path

sys.stdout.reconfigure(encoding="utf-8")
sys.stderr.reconfigure(encoding="utf-8")

sys.path.insert(0, str(Path(__file__).resolve().parent))
from wechat_export.runtime_paths import validate_runtime_directory  # noqa: E402

INDEX_HEADERS = [
    "序号", "客户显示名", "备注名", "昵称", "微信号", "微信号别名",
    "消息数", "我发送", "客户发送", "未知方向", "首条时间", "末条时间",
    "HTML文件", "CSV文件",
]

_COLUMN_WIDTHS = {
    "A": 8,
    "B": 22, "C": 22, "D": 22, "E": 22, "F": 22,
    "G": 12, "H": 12, "I": 12, "J": 12,
    "K": 20, "L": 20,
    "M": 44, "N": 44,
}

_INTEGER_COLUMNS = ("A",)
_COUNT_COLUMNS = ("G", "H", "I", "J")


def _spreadsheet_safe(value: object) -> object:
    if isinstance(value, str) and value.startswith(("=", "+", "-", "@", "\t", "\r")):
        return "'" + value
    return value


def _write_state(run_root: Path, payload: dict[str, object]) -> None:
    (run_root / "run-state.json").write_text(
        json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8"
    )


def _read_state(run_root: Path) -> dict[str, object]:
    state_path = run_root / "run-state.json"
    if state_path.is_file():
        return json.loads(state_path.read_text(encoding="utf-8"))
    return {}


def _build_xlsx(source: dict, rows: list[dict], out_path: Path) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.table import Table, TableStyleInfo

    header_fill = PatternFill("solid", fgColor="17324D")
    header_font = Font(bold=True, color="FFFFFF")
    title_font = Font(bold=True, color="FFFFFF", size=18)

    workbook = Workbook()
    summary = workbook.active
    summary.title = "汇总"
    index_sheet = workbook.create_sheet("客户索引")

    # 汇总
    summary.sheet_view.showGridLines = False
    summary.merge_cells("A1:F1")
    summary["A1"] = "微信客户一对一聊天记录总览"
    summary["A1"].fill = header_fill
    summary["A1"].font = title_font
    summary["A1"].alignment = Alignment(horizontal="left", vertical="center")
    summary.row_dimensions[1].height = 38

    summary.merge_cells("A2:F2")
    summary["A2"] = (
        "范围：已勾选的一对一客户会话；已排除群聊、公众号、系统账号与未勾选会话。"
        "数据仅在本机处理。"
    )
    summary["A2"].alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
    summary.row_dimensions[2].height = 30

    direction_totals = source.get("direction_totals") or {}
    metrics = [
        ("指标", "数值"),
        ("客户会话数", len(rows)),
        ("一对一消息总数", source.get("private_messages", 0)),
        ("我发送", direction_totals.get("我", 0)),
        ("客户发送", direction_totals.get("客户", 0)),
        ("未知方向", direction_totals.get("未知", 0)),
        ("已筛除会话数", source.get("screened_out_conversations", 0)),
        ("已筛除消息数", source.get("screened_out_messages", 0)),
    ]
    for offset, (label, value) in enumerate(metrics):
        row = 4 + offset
        summary.cell(row=row, column=1, value=label)
        summary.cell(row=row, column=2, value=value)
    summary["A4"].fill = header_fill
    summary["A4"].font = header_font
    summary["B4"].fill = header_fill
    summary["B4"].font = header_font
    for offset in range(1, len(metrics)):
        row = 4 + offset
        cell = summary.cell(row=row, column=2)
        cell.number_format = "#,##0"
        cell.font = Font(bold=True, color="17324D")
        cell.fill = PatternFill("solid", fgColor="F7FAFC")
    summary.column_dimensions["A"].width = 24
    summary.column_dimensions["B"].width = 18

    # 客户索引
    index_sheet.sheet_view.showGridLines = False
    index_sheet.append(INDEX_HEADERS)
    for cell in index_sheet[1]:
        cell.fill = header_fill
        cell.font = header_font
    index_sheet.row_dimensions[1].height = 28
    for row in rows:
        index_sheet.append([_spreadsheet_safe(row.get(key, "")) for key in INDEX_HEADERS])

    last_row = len(rows) + 1
    if last_row >= 2:
        last_col = get_column_letter(len(INDEX_HEADERS))
        table = Table(displayName="CustomerChatIndex", ref=f"A1:{last_col}{last_row}")
        table.tableStyleInfo = TableStyleInfo(
            name="TableStyleMedium2", showRowStripes=True
        )
        index_sheet.add_table(table)

    for column, width in _COLUMN_WIDTHS.items():
        index_sheet.column_dimensions[column].width = width
    for column in _INTEGER_COLUMNS:
        for row in range(2, last_row + 1):
            index_sheet[f"{column}{row}"].number_format = "0"
    for column in _COUNT_COLUMNS:
        for row in range(2, last_row + 1):
            index_sheet[f"{column}{row}"].number_format = "#,##0"

    index_sheet.freeze_panes = "C2"
    index_sheet.auto_filter.ref = f"A1:{get_column_letter(len(INDEX_HEADERS))}{last_row}"

    out_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(out_path)


def _write_csv_fallback(rows: list[dict], out_path: Path) -> None:
    out_path.parent.mkdir(parents=True, exist_ok=True)
    with out_path.open("w", encoding="utf-8-sig", newline="") as stream:
        writer = csv.writer(stream)
        writer.writerow(INDEX_HEADERS)
        for row in rows:
            writer.writerow([_spreadsheet_safe(row.get(key, "")) for key in INDEX_HEADERS])


def build_index(run_root: Path) -> dict[str, object]:
    run_root = validate_runtime_directory(run_root)
    delivery = run_root / "delivery"
    index_json = delivery / "客户索引数据.json"
    if not index_json.is_file():
        raise FileNotFoundError(
            f"找不到客户索引数据：{index_json}；请先运行 run_export.py 完成导出"
        )
    source = json.loads(index_json.read_text(encoding="utf-8"))
    rows = source.get("index_rows", [])

    state = _read_state(run_root)
    try:
        import openpyxl  # noqa: F401
    except ImportError:
        csv_path = delivery / "客户聊天索引.csv"
        _write_csv_fallback(rows, csv_path)
        result = {
            "status": "csv_fallback",
            "index_path": str(csv_path),
            "rows": len(rows),
            "next_action": (
                "未安装 openpyxl，已降级为 CSV 索引；如需 Excel 请征得用户同意后 "
                "pip install openpyxl 重跑本步"
            ),
        }
    else:
        xlsx_path = delivery / "客户聊天索引.xlsx"
        _build_xlsx(source, rows, xlsx_path)
        result = {
            "status": "ok",
            "index_path": str(xlsx_path),
            "rows": len(rows),
        }

    phases = list(state.get("phases", []))
    phases.append({"phase": "build_index", "status": "ok", "rows": len(rows)})
    state["phases"] = phases
    state["status"] = "needs_finalize"
    state["index_path"] = result["index_path"]
    _write_state(run_root, state)
    return result


def main() -> int:
    parser = argparse.ArgumentParser(
        description="Build the customer chat index spreadsheet (Excel via openpyxl, "
        "or CSV fallback if openpyxl is unavailable)"
    )
    parser.add_argument("--run-root", required=True, type=Path)
    args = parser.parse_args()

    run_root = validate_runtime_directory(args.run_root)
    try:
        result = build_index(run_root)
        print(json.dumps(result, ensure_ascii=False, sort_keys=True))
        return 0
    except Exception as exc:
        state = _read_state(run_root)
        state["status"] = "failed"
        state["error_type"] = type(exc).__name__
        state["error"] = str(exc)
        _write_state(run_root, state)
        print(
            json.dumps(
                {
                    "status": "failed",
                    "run_root": str(run_root),
                    "error_type": type(exc).__name__,
                    "error": str(exc),
                },
                ensure_ascii=False,
                sort_keys=True,
            ),
            file=sys.stderr,
        )
        return 1


if __name__ == "__main__":
    raise SystemExit(main())
